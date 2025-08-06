# main.py - Intelligent Claims Adjudicator Backend

import os
import json
import fitz
import docx
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import HuggingFaceEmbeddings
import google.generativeai as genai
from flask import Flask, request, jsonify, render_template, url_for
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

# Import the history manager module
import history_manager

# --- FLASK APP INITIALIZATION & SECURITY ---
# Make the UPLOAD_FOLDER accessible for the frontend to fetch documents
app = Flask(__name__, static_folder='policy_documents')
load_dotenv()

# CRITICAL SECURITY: Load API key from environment variables.
API_KEY = os.getenv("GOOGLE_API_KEY")
if not API_KEY:
    print("WARNING: GOOGLE_API_KEY not found in environment variables. Using a placeholder.")
    API_KEY = "YOUR_API_KEY_HERE" # Replace with your actual key if not using .env

genai.configure(api_key=API_KEY)

# --- CONFIGURATION ---
UPLOAD_FOLDER = "policy_documents"
VECTOR_STORE_DIR = "faiss_vector_stores"
EMBEDDING_MODEL_NAME = "all-MiniLM-L6-v2"
ALLOWED_EXTENSIONS = {'txt', 'pdf', 'docx'}
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(VECTOR_STORE_DIR, exist_ok=True)

# --- GLOBAL EMBEDDING MODEL ---
print("Initializing embedding model...")
embeddings = HuggingFaceEmbeddings(model_name=EMBEDDING_MODEL_NAME)
print("Embedding model loaded.")

# --- DOCUMENT PROCESSING & VECTOR STORE LOGIC ---
# (These functions remain unchanged)
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_file(file_path, filename):
    file_text = ""
    try:
        if filename.endswith(".pdf"):
            with fitz.open(file_path) as doc:
                # Store text per page for metadata
                pages_text = []
                for page_num, page in enumerate(doc):
                    pages_text.append({"text": page.get_text(), "metadata": {"page": page_num + 1}})
                return pages_text
        elif filename.endswith(".txt"):
            with open(file_path, 'r', encoding='utf-8') as f:
                # For txt, consider the whole file as page 1
                return [{"text": f.read(), "metadata": {"page": 1}}]
        elif filename.endswith(".docx"):
            doc = docx.Document(file_path)
            # For docx, consider the whole file as page 1
            full_text = "\n".join([para.text for para in doc.paragraphs])
            return [{"text": full_text, "metadata": {"page": 1}}]
    except Exception as e:
        print(f"Error reading file {filename}: {e}")
        return None

def create_and_save_vector_store(filename):
    file_path = os.path.join(UPLOAD_FOLDER, filename)
    store_path = os.path.join(VECTOR_STORE_DIR, f"{filename}.faiss")
    if os.path.exists(store_path): return True
    
    pages = extract_text_from_file(file_path, filename)
    if not pages: return False

    # Use LangChain's Document object to store metadata
    from langchain_core.documents import Document
    all_docs = [Document(page_content=page['text'], metadata=page['metadata']) for page in pages]

    text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)
    chunks = text_splitter.split_documents(all_docs)
    if not chunks: return False

    store = FAISS.from_documents(documents=chunks, embedding=embeddings)
    store.save_local(store_path)
    return True

def get_retriever(selected_policy):
    store_path = os.path.join(VECTOR_STORE_DIR, f"{selected_policy}.faiss")
    if not os.path.exists(store_path):
        if not create_and_save_vector_store(selected_policy): return None
    try:
        vector_store = FAISS.load_local(store_path, embeddings, allow_dangerous_deserialization=True)
        return vector_store.as_retriever(search_kwargs={"k": 7})
    except Exception as e:
        print(f"Error loading vector store {store_path}: {e}")
        return None

# --- INTELLIGENCE CORE ---
def get_adjudication_from_llm(chat_history: list, relevant_docs: list, user_query: str):
    model = genai.GenerativeModel('gemini-2.5-flash-lite')
    
    # NEW: Format the context to include page numbers
    clauses_with_metadata = [
        {"content": doc.page_content, "page": doc.metadata.get('page', 1)} 
        for doc in relevant_docs
    ]

    prompt = f"""
    You are a strict, senior insurance claim adjudicator AI. Your primary directive is to provide a definitive, evidence-based decision based on the provided policy clauses.

    **Analysis Context:**
    - **User's Query:** "{user_query}"
    - **Relevant Policy Clauses (Source of Truth):**
    ---
    {json.dumps(clauses_with_metadata, indent=2)}
    ---

    **Your Mandate:**
    1.  **Analyze:** Evaluate the user's query against the rules in the "Relevant Policy Clauses" JSON above.
    2.  **Decide:** Determine the final claim status: "Approved" or "Rejected".
    3.  **Determine Payout Amount:** If Approved, you MUST find the specific monetary value from the clauses.
    4.  **Justify:** Write a comprehensive justification for your decision.
    5.  **Cite Evidence:** This is critical. Your response for "cited_clauses" MUST be a JSON array of objects. Each object must contain the exact text of the clause you are citing (`text`) and the corresponding page number (`page`) from the source JSON.

    **Required Output Format:**
    You MUST respond with a single, valid JSON object.
    {{
      "decision": "string (Must be 'Approved' or 'Rejected')",
      "amount": "string (The calculated payout amount, e.g., '₹75,000', or null if Rejected)",
      "justification": "string (Your detailed explanation).",
      "cited_clauses": [
        {{
          "text": "string (Direct quote of the supporting clause)",
          "page": "integer (The page number from the source JSON)"
        }}
      ]
    }}
    """
    try:
        response = model.generate_content(
            prompt,
            generation_config={"response_mime_type": "application/json"}
        )
        return json.loads(response.text)
    except Exception as e:
        print(f"Error processing adjudication with LLM: {e}")
        return {
            "decision": "Error", "amount": None,
            "justification": "I encountered an internal error. Please try rephrasing your question.",
            "cited_clauses": []
        }

# --- WEB ROUTES ---
@app.route('/')
def home():
    return render_template('index.html')

# NEW: Add a route to serve the documents
@app.route('/documents/<path:filename>')
def serve_document(filename):
    return app.send_static_file(filename)

@app.route('/upload_policy', methods=['POST'])
def upload_policy():
    if 'policyFile' not in request.files: return jsonify({'error': 'No file part.'}), 400
    file = request.files['policyFile']
    if file.filename == '' or not allowed_file(file.filename): return jsonify({'error': 'Invalid file type.'}), 400
    filename = secure_filename(file.filename)
    save_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    try:
        file.save(save_path)
        if create_and_save_vector_store(filename):
            return jsonify({'message': f'Successfully indexed {filename}', 'filename': filename}), 200
        else:
            os.remove(save_path)
            return jsonify({'error': f'The uploaded file "{filename}" appears to be empty or could not be read.'}), 500
    except Exception as e:
        return jsonify({'error': f'Server error during file upload: {e}'}), 500

@app.route('/list_policies')
def list_policies():
    try:
        policies = [f for f in os.listdir(UPLOAD_FOLDER) if allowed_file(f)]
        return jsonify({'policies': policies})
    except Exception as e:
        return jsonify({'error': f'Could not list policies: {e}'}), 500

@app.route('/chat', methods=['POST'])
def chat_route():
    data = request.get_json()
    user_query, policy, history = data.get('query'), data.get('policy'), data.get('history', [])
    if not all([user_query, policy]):
        return jsonify({'error': 'A query and a selected policy are required.'}), 400

    retriever = get_retriever(policy)
    if not retriever:
        return jsonify({"error": "The selected policy document could not be loaded or processed."}), 500

    # Retrieve documents with their metadata
    retrieved_docs = retriever.invoke(user_query)
    
    # Pass the full document objects to the LLM function
    bot_response_json = get_adjudication_from_llm(history, retrieved_docs, user_query)
    
    # Add the document URL to the response for the frontend
    bot_response_json['document_url'] = url_for('serve_document', filename=policy, _external=True)
    
    history_manager.add_claim_to_history(user_query, policy, bot_response_json)
    return jsonify(bot_response_json)

@app.route('/get_history', methods=['GET'])
def get_history():
    claim_history = history_manager.read_history()
    return jsonify(claim_history)

# --- MAIN EXECUTION ---
if __name__ == "__main__":
    print("--- Starting Intelligent Policy Adjudicator Server ---")
    print("Indexing any existing documents...")
    for filename in os.listdir(UPLOAD_FOLDER):
        if allowed_file(filename):
            create_and_save_vector_store(filename)
    print(f"--- Server ready. Access at http://127.0.0.1:5000 ---")
    app.run(host='0.0.0.0', port=5000, debug=True)
